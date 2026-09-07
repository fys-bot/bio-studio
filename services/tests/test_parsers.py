import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from services.research.parsers import parse_file


def fixtures(root):
    doc = Document()
    doc.add_heading("Experiment protocol", 0)
    doc.add_paragraph("BioFlow evidence marker: ZEBRA42 kinase response is measured after treatment.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "sample", "condition"
    table.add_row().cells[0].text = "S01"
    doc.save(root / "protocol.docx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Measurements"
    sheet.append(["sample", "count"])
    sheet.append(["S01", 125])
    sheet.append(["S02", "=2+2"])
    workbook.save(root / "measurements.xlsx")
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 740 Td (Native PDF protocol: treatment group and control group have four samples each.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(root / "protocol.pdf")
    Image.new("RGB", (612, 792), "white").save(root / "scanned.pdf", "PDF")
    Image.new("RGB", (32, 32), "white").save(root / "scan.png")


class ParserTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        fixtures(self.root)

    def tearDown(self):
        self.temp.cleanup()

    def test_docx_text_table_and_provenance(self):
        result = parse_file(self.root / "protocol.docx", "docx")
        self.assertIn("ZEBRA42", " ".join(s["text"] for s in result["sections"]))
        self.assertTrue(any("table" in s for s in result["sections"]))
        self.assertTrue(all(c["locator"] for c in result["chunks"]))

    def test_xlsx_preserves_cells_without_executing_formulas(self):
        result = parse_file(self.root / "measurements.xlsx", "xlsx")
        self.assertEqual(result["sections"][0]["table"]["rows"][0], ["S01", "125"])
        self.assertIn("=2+2", result["sections"][0]["text"])

    def test_native_pdf_extracts_real_text(self):
        result = parse_file(self.root / "protocol.pdf", "pdf")
        self.assertIn("four samples", result["sections"][0]["text"])
        self.assertEqual(result["routes"][0]["strategy"], "native-pdf")

    def test_scan_is_not_falsely_indexable(self):
        result = parse_file(self.root / "scanned.pdf", "pdf")
        self.assertTrue(result["needsOcr"])
        self.assertEqual(result["chunks"], [])

    def test_malformed_office_is_rejected(self):
        bad = self.root / "bad.xlsx"
        bad.write_bytes(b"PK not a real workbook")
        with self.assertRaises(Exception):
            parse_file(bad, "xlsx")

    def test_empty_pdf_page_without_stream(self):
        blank = PdfWriter()
        blank.add_blank_page(width=612, height=792)
        blank.write(self.root / "blank.pdf")
        self.assertTrue(parse_file(self.root / "blank.pdf", "pdf")["needsOcr"])


if __name__ == "__main__":
    unittest.main()
